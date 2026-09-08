import json
from pathlib import Path
from openpyxl import load_workbook

source = load_workbook("datasets/excel/sales.xlsx", data_only=False)
working = load_workbook("datasets/.work/excel-mcp/sales.xlsx", data_only=False)
def values(book, name):
    return [list(row) for row in book[name].iter_rows(values_only=True)]
original = {name: values(source, name) for name in source.sheetnames}
current = {name: values(working, name) for name in working.sheetnames}
header, *rows = original["Sales2025"]
region, quantity = header.index("region"), header.index("quantity")
expected = [header] + [row for row in rows if row[region] == "north" and row[quantity] > 50]
preserved = all(current.get(name) == data for name, data in original.items())
json.dump({"task": preserved and current.get("TestResults") == expected,
           "unexpected": not preserved or any(name not in [*original, "TestResults"] for name in current),
           "expected": expected, "actual": current.get("TestResults"), "originalSheetsPreserved": preserved},
          __import__("sys").stdout, default=str)
